import pandas as pd
from openpyxl import load_workbook
from tkinter import filedialog

class FileHandler:
    @staticmethod
    def load_excel_file(app):
        """Otvori dijalog za odabir file i učitaj Excel"""
        try:
            filepath = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
            if not filepath:
                return None, None
            
            workbook = load_workbook(filepath)
            sheet_names = workbook.sheetnames
            app.update_sheets(sheet_names, filepath)  # Ažuriranje GUI-ja
            
            df = pd.read_excel(filepath)
            return df, filepath

        except Exception as e:
            app.display_message(f"Greška: {e}")
            return None, None
